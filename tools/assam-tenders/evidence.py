"""Offline, deterministic evidence retention. Never executes workbook formulas/macros."""

import hashlib
import html
import json
import subprocess
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlsplit
from assam_tenders import Node

VERSION = 1


def safe_route(url):
    parsed = urlsplit(url)
    query = parse_qs(parsed.query)
    return parsed.path + (
        "?"
        + urlencode(
            {k: query[k][0] for k in ("page", "component", "service") if k in query}
        )
        if query
        else ""
    )


def sanitized_html(node):
    if not isinstance(node, Node):
        return html.escape(node)
    if node.tag in {"script", "style", "input", "textarea", "select", "img"}:
        return ""
    attrs = {
        k: v
        for k, v in node.attrs.items()
        if k in {"id", "class", "title", "colspan", "rowspan"}
    }
    if node.tag == "a" and node.attrs.get("href", "").startswith(
        ("/", "https://assamtenders.gov.in/")
    ):
        attrs["href"] = safe_route(node.attrs["href"])
    content = "".join(sanitized_html(c) for c in node.children)
    if node.tag == "root":
        return content
    return (
        "<"
        + node.tag
        + "".join(
            " " + k + '="' + html.escape(v or "", quote=True) + '"'
            for k, v in attrs.items()
        )
        + ">"
        + content
        + "</"
        + node.tag
        + ">"
    )


def snapshot(page):
    dom = page.dom
    return {
        "schema_version": VERSION,
        "route": safe_route(page.url),
        "html": sanitized_html(dom),
        "text": dom.text(),
        "tables": [
            [
                [
                    c.text()
                    for c in row.children
                    if isinstance(c, Node) and c.tag in {"td", "th"}
                ]
                for row in table.find_all("tr")
            ]
            for table in dom.find_all("table")
        ],
    }


def retain_snapshot(root, tender_id, kind, page):
    data = snapshot(page)
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True).encode()
    digest = hashlib.sha256(payload).hexdigest()
    folder = Path(root) / "evidence" / tender_id / kind
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / (digest + ".json")
    if not target.exists():
        target.write_bytes(payload)
    return str(target.relative_to(root))


def extract(path, root):
    """Preserve all PDF pages and populated workbook cells; originals remain authoritative."""
    path, root = Path(path), Path(root)
    hasher = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1048576), b""):
            hasher.update(chunk)
    digest = hasher.hexdigest()
    folder = root / "extracted" / digest
    folder.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    result = {"extractor_version": VERSION, "source_sha256": digest}
    if suffix == ".pdf":
        target = folder / "text.txt"
        proc = subprocess.run(
            ["pdftotext", "-layout", "-enc", "UTF-8", str(path), str(target)],
            capture_output=True,
            timeout=120,
        )
        if proc.returncode:
            raise ValueError("PDF text extraction failed: " + path.name)
        text = target.read_text()
        result.update(
            status="text" if text.strip() else "no_text_layer",
            characters=len(text),
            text_path=str(target.relative_to(root)),
        )
    elif suffix in {".xls", ".xlsx"}:
        sheets = []
        if suffix == ".xls":
            import xlrd

            book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
            try:
                for sheet in book.sheets():
                    if sheet.nrows * sheet.ncols > 2_000_000:
                        raise ValueError(
                            "Workbook sheet exceeds two million cell extraction limit; original retained."
                        )
                    cells = [
                        {
                            "row": r + 1,
                            "column": c + 1,
                            "type": sheet.cell_type(r, c),
                            "value": sheet.cell_value(r, c),
                        }
                        for r in range(sheet.nrows)
                        for c in range(sheet.ncols)
                        if sheet.cell_type(r, c) not in {0, 6}
                    ]
                    sheets.append(
                        {
                            "name": sheet.name,
                            "rows": sheet.nrows,
                            "columns": sheet.ncols,
                            "visibility": sheet.visibility,
                            "merged_cells": sheet.merged_cells,
                            "cells": cells,
                        }
                    )
                result["excel_date_mode"] = book.datemode
                result["formula_note"] = (
                    "XLS values are cached results; original binary retains formulas and macros."
                )
            finally:
                book.release_resources()
        else:
            import openpyxl

            book = openpyxl.load_workbook(
                path, read_only=True, data_only=False, keep_links=False
            )
            try:
                for sheet in book:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 2_000_000:
                        raise ValueError(
                            "Workbook sheet exceeds extraction limit; original retained."
                        )
                    cells = [
                        {
                            "row": cell.row,
                            "column": cell.column,
                            "type": cell.data_type,
                            "value": (
                                str(cell.value)
                                if not isinstance(cell.value, (str, int, float, bool))
                                else cell.value
                            ),
                        }
                        for row in sheet
                        for cell in row
                        if cell.value is not None
                    ]
                    sheets.append(
                        {
                            "name": sheet.title,
                            "visibility": sheet.sheet_state,
                            "cells": cells,
                        }
                    )
            finally:
                book.close()
        target = folder / "workbook.json"
        target.write_text(json.dumps({"sheets": sheets}, ensure_ascii=False, indent=2))
        result.update(
            status="workbook",
            sheets=len(sheets),
            cells=sum(len(s["cells"]) for s in sheets),
            workbook_path=str(target.relative_to(root)),
        )
    else:
        result["status"] = "original_preserved_unsupported_extraction"
    target = folder / "extraction.json"
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    return result
